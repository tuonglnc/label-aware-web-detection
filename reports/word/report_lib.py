"""Helpers to build the report on top of the faculty Word template."""
import copy
import re

from lxml import etree
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

XSL = etree.XSLT(etree.parse(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))
BODY_STYLE = "Nội dung văn bản"
TEXT_WIDTH_CM = 15.5


# ---------- MathML helpers ----------
def mi(x): return f"<mi>{x}</mi>"
def mn(x): return f"<mn>{x}</mn>"
def mo(x): return f"<mo>{x}</mo>"
def mt(x): return f"<mtext>{x}</mtext>"
def row(*a): return "<mrow>" + "".join(a) + "</mrow>"
def frac(a, b): return f"<mfrac>{row(a)}{row(b)}</mfrac>"
def sqrt(a): return f"<msqrt>{row(a)}</msqrt>"
def sub(a, b): return f"<msub>{row(a)}{row(b)}</msub>"
def sup(a, b): return f"<msup>{row(a)}{row(b)}</msup>"
def subsup(a, b, c): return f"<msubsup>{row(a)}{row(b)}{row(c)}</msubsup>"
def bar(a): return f"<mover accent='true'>{row(a)}<mo>&#x00AF;</mo></mover>"
def hat(a): return f"<mover accent='true'>{row(a)}<mo>^</mo></mover>"
def paren(*a): return "<mfenced>" + row(*a) + "</mfenced>"
def summ(lo, hi, body): return f"<munderover><mo>&#x2211;</mo>{row(lo)}{row(hi)}</munderover>{row(body)}"
def sumu(lo, body): return f"<munder><mo>&#x2211;</mo>{row(lo)}</munder>{row(body)}"


def omml(mathml):
    """Convert a MathML fragment to a Word math element."""
    src = f'<math xmlns="http://www.w3.org/1998/Math/MathML">{mathml}</math>'
    return parse_xml(etree.tostring(XSL(etree.fromstring(src))))


TOKEN = re.compile(r"(\[/?(?:b|i|sub|sup|code)\]|\[m:[a-z0-9_]+\])")


class Report:
    def __init__(self, template_docx, inline_math=None):
        self.doc = Document(template_docx)
        self.body = self.doc.element.body
        self.inline = inline_math or {}
        self.chapter = 0
        self.count = {"Hình": 0, "Bảng": 0, "Equation": 0}
        self.refs = {}
        self.eq_template = None

    # ---------- low level ----------
    def children(self):
        return list(self.body.iterchildren())

    def runs(self, par, text, size=None, bold=None, italic=None, font=None):
        state = {"b": False, "i": False, "sub": False, "sup": False, "code": False}
        for part in TOKEN.split(text):
            if not part:
                continue
            m = re.fullmatch(r"\[(/?)(b|i|sub|sup|code)\]", part)
            if m:
                state[m.group(2)] = m.group(1) == ""
                continue
            m = re.fullmatch(r"\[m:([a-z0-9_]+)\]", part)
            if m:
                par._p.append(omml(self.inline[m.group(1)]))
                continue
            run = par.add_run(part)
            if size:
                run.font.size = Pt(size)
            if bold or state["b"]:
                run.bold = True
            if italic or state["i"]:
                run.italic = True
            if state["sub"]:
                run.font.subscript = True
            if state["sup"]:
                run.font.superscript = True
            if state["code"] or font:
                name = font or "Consolas"
                run.font.name = name
                rpr = run._r.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                rfonts.set(qn("w:hAnsi"), name)
                rfonts.set(qn("w:cs"), name)
                if state["code"] and not size:
                    run.font.size = Pt(11.5)
        return par

    # ---------- headings and text ----------
    def h1(self, text, numbered=True, page_break=True):
        par = self.doc.add_paragraph(style="Heading 1")
        par.paragraph_format.page_break_before = page_break
        par.paragraph_format.keep_with_next = True
        if numbered:
            self.chapter += 1
            self.count = {k: 0 for k in self.count}
        else:
            ppr = par._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
            numid = OxmlElement("w:numId"); numid.set(qn("w:val"), "0")
            numpr.append(ilvl); numpr.append(numid)
            ppr.append(numpr)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run(text)
        return par

    def h2(self, text):
        par = self.doc.add_paragraph(style="Heading 2")
        par.paragraph_format.keep_with_next = True
        par.add_run(text)
        return par

    def h3(self, text):
        par = self.doc.add_paragraph(style="Heading 3")
        par.paragraph_format.keep_with_next = True
        par.add_run(text)
        return par

    def p(self, text, indent=True, align=None, keep=False, bold=None, italic=None, left_cm=None):
        par = self.doc.add_paragraph(style=BODY_STYLE)
        if not indent:
            par.paragraph_format.first_line_indent = Cm(0)
        if left_cm is not None:
            par.paragraph_format.left_indent = Cm(left_cm)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if keep:
            par.paragraph_format.keep_with_next = True
        return self.runs(par, text, bold=bold, italic=italic)

    def bullets(self, items):
        for item in items:
            par = self.doc.add_paragraph(style=BODY_STYLE)
            par.paragraph_format.left_indent = Cm(1.27)
            par.paragraph_format.first_line_indent = Cm(-0.5)
            self.runs(par, "- " + item)

    def gap(self, pt=6):
        par = self.doc.add_paragraph(style=BODY_STYLE)
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.line_spacing = 1.0
        run = par.add_run("")
        run.font.size = Pt(pt)
        ppr = par._p.get_or_add_pPr()
        rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(pt * 2)))
        rpr.append(sz); ppr.append(rpr)

    # ---------- captions and references ----------
    def _field(self, par, instr, value):
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), instr)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t"); t.text = value
        r.append(t); fld.append(r)
        par._p.append(fld)

    def caption(self, kind, key, text, keep=False):
        self.count[kind] += 1
        number = f"{self.chapter}.{self.count[kind]}"
        self.refs[key] = f"{kind} {number}"
        par = self.doc.add_paragraph(style="Caption")
        par.paragraph_format.keep_with_next = keep
        par.add_run(f"{kind} ")
        self._field(par, " STYLEREF 1 \\s ", str(self.chapter))
        par.add_run(".")
        self._field(par, f" SEQ {kind} \\* ARABIC \\s 1 ", str(self.count[kind]))
        par.add_run(": " + text)
        return par

    def fig(self, key, path, caption, width_cm=TEXT_WIDTH_CM):
        par = self.doc.add_paragraph(style=BODY_STYLE)
        par.paragraph_format.first_line_indent = Cm(0)
        par.paragraph_format.line_spacing = 1.0
        par.paragraph_format.space_before = Pt(6)
        par.paragraph_format.keep_with_next = True
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        self.caption("Hình", key, caption)

    # ---------- tables ----------
    def _borders(self, table, color="000000", size="4"):
        tblpr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), size)
            el.set(qn("w:space"), "0"); el.set(qn("w:color"), color)
            borders.append(el)
        tblpr.append(borders)

    def _cell(self, cell, text, size, bold=False, align="left", font=None):
        par = cell.paragraphs[0]
        par.style = self.doc.styles["Normal"]
        fmt = par.paragraph_format
        fmt.space_before = Pt(2); fmt.space_after = Pt(2); fmt.line_spacing = 1.0
        par.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                         "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        self.runs(par, str(text), size=size, bold=bold, font=font)

    def table(self, key, caption, header, rows, widths, size=12, aligns=None):
        self.caption("Bảng", key, caption, keep=True)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._borders(table)
        aligns = aligns or ["left"] + ["center"] * (len(header) - 1)
        for j, text in enumerate(header):
            self._cell(table.cell(0, j), text, size, bold=True, align="center")
        for i, values in enumerate(rows, start=1):
            for j, text in enumerate(values):
                self._cell(table.cell(i, j), text, size, align=aligns[j])
        for table_row in table.rows:
            trpr = table_row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            for j, cell in enumerate(table_row.cells):
                cell.width = Cm(widths[j])
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        header_pr.append(OxmlElement("w:tblHeader"))
        if len(table.rows) <= 25:
            for table_row in table.rows[:-1]:
                for cell in table_row.cells:
                    for par in cell.paragraphs:
                        par.paragraph_format.keep_with_next = True
        self.gap()
        return table

    def code(self, text, size=8.5):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._borders(table, color="BFBFBF")
        cell = table.cell(0, 0)
        cell.width = Cm(TEXT_WIDTH_CM)
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2")
        tcpr.append(shd)
        lines = text.strip("\n").split("\n")
        for k, line in enumerate(lines):
            par = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            par.style = self.doc.styles["Normal"]
            fmt = par.paragraph_format
            fmt.space_before = Pt(0); fmt.space_after = Pt(0); fmt.line_spacing = 1.0
            run = par.add_run(line if line else " ")
            run.font.name = "Consolas"; run.font.size = Pt(size)
            rfonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
            rfonts.set(qn("w:hAnsi"), "Consolas"); rfonts.set(qn("w:cs"), "Consolas")
        self.gap()

    # ---------- equations ----------
    def eq(self, key, mathml):
        self.count["Equation"] += 1
        number = f"{self.chapter}.{self.count['Equation']}"
        self.refs[key] = f"({number})"
        tbl = copy.deepcopy(self.eq_template)
        for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
            for el in tbl.findall(".//" + qn(tag)):
                el.getparent().remove(el)
        old = tbl.find(".//" + qn("m:oMathPara"))
        para = OxmlElement("m:oMathPara")
        para.append(omml(mathml))
        old.getparent().replace(old, para)
        cached = [self.chapter, self.count["Equation"]]
        for fld, value in zip(tbl.findall(".//" + qn("w:fldSimple")), cached):
            fld.find(".//" + qn("w:t")).text = str(value)
        self.body.insert(len(self.body) - 1, tbl)

    # ---------- finishing ----------
    def resolve_refs(self):
        pattern = re.compile(r"<<([a-z0-9:_]+)>>")
        missing = set()
        for t in self.body.iter(qn("w:t")):
            if t.text and "<<" in t.text:
                def repl(m):
                    if m.group(1) not in self.refs:
                        missing.add(m.group(1))
                        return m.group(0)
                    return self.refs[m.group(1)]
                t.text = pattern.sub(repl, t.text)
        return missing


def set_text(p_el, text):
    """Put the text in the first text node of a paragraph and blank the rest."""
    nodes = p_el.findall(".//" + qn("w:t"))
    nodes[0].text = text
    nodes[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for node in nodes[1:]:
        node.text = ""


def rebuild(p_el, parts, rpr_index=0):
    """Replace all runs of a paragraph by new runs: parts is a list of (text, bold)."""
    runs = p_el.findall(qn("w:r"))
    base = copy.deepcopy(runs[min(rpr_index, len(runs) - 1)].find(qn("w:rPr")))
    for tag in ("w:highlight", "w:b", "w:bCs", "w:spacing"):
        for el in base.findall(qn(tag)):
            base.remove(el)
    for r in runs:
        p_el.remove(r)
    for text, bold in parts:
        r = OxmlElement("w:r")
        rpr = copy.deepcopy(base)
        if bold:
            rpr.insert(0, OxmlElement("w:b"))
        r.append(rpr)
        t = OxmlElement("w:t"); t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p_el.append(r)
